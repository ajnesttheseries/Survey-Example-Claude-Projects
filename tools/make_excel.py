# -*- coding: utf-8 -*-
"""สร้างไฟล์ Excel แคตตาล็อกผลงาน AI จาก data/catalog.json"""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

catalog = json.load(open('data/catalog.json'))
POST = 'https://www.facebook.com/groups/1745892855948687/posts/2323047211566579/'

NAVY   = '1F3864'
BAND   = 'EAF1F8'
AMBER  = 'FFF2CC'
GREY   = '808080'
BLUE   = '0563C1'

F  = 'Arial'
thin = Side(style='thin', color='BFBFBF')
box  = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ─────────────────────────── ชีต 1: แคตตาล็อก ───────────────────────────
ws = wb.active
ws.title = 'แคตตาล็อกผลงาน'

HEAD = ['ลำดับ', 'หมวดหมู่', 'ชื่อผลงาน', 'รายละเอียด', 'ผู้สร้าง',
        'ลิงก์ผลงาน', 'ลิงก์เพิ่มเติม', 'ยอดรีแอ็กชัน', 'สถานะลิงก์', 'คอมเมนต์ต้นทาง']
WIDTH = [7, 26, 30, 62, 26, 40, 34, 12, 30, 16]

ws['A1'] = 'แคตตาล็อกผลงานที่สร้างด้วย AI'
ws['A1'].font = Font(name=F, size=15, bold=True, color=NAVY)
ws['A2'] = f'รวบรวมจากคอมเมนต์ในโพสต์ Facebook · {len(catalog)} ผลงาน · 16 หมวดหมู่'
ws['A2'].font = Font(name=F, size=10, color=GREY)
ws['A3'] = 'ที่มาของโพสต์'
ws['A3'].font = Font(name=F, size=10, color=BLUE, underline='single')
ws['A3'].hyperlink = POST

HEADER_ROW = 5
for c, (h, w) in enumerate(zip(HEAD, WIDTH), 1):
    cell = ws.cell(row=HEADER_ROW, column=c, value=h)
    cell.font = Font(name=F, size=10, bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor=NAVY)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = box
    ws.column_dimensions[get_column_letter(c)].width = w
ws.row_dimensions[HEADER_ROW].height = 26

for i, r in enumerate(catalog):
    row = HEADER_ROW + 1 + i
    ok = r['linkStatus'] == 'ใช้ได้'

    vals = [r['no'], r['catName'], r['name'], r['desc'], r['author'],
            r['url'], '\n'.join(r['extras']), r['reactions'],
            'ใช้ได้' if ok else 'ลิงก์ถูก Facebook ตัด — ต้องค้นหาเอง', 'เปิดคอมเมนต์']
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = Font(name=F, size=10)
        cell.border = box
        cell.alignment = Alignment(vertical='top', wrap_text=(c in (2, 3, 4, 7, 9)))

    ws.cell(row=row, column=1).alignment = Alignment(horizontal='center', vertical='top')
    ws.cell(row=row, column=8).alignment = Alignment(horizontal='center', vertical='top')
    ws.cell(row=row, column=8).number_format = '#,##0'

    # ลิงก์ผลงาน — ทำเป็น hyperlink เฉพาะอันที่กดแล้วไปถึงจริง
    lc = ws.cell(row=row, column=6)
    if ok:
        lc.hyperlink = r['url']
        lc.font = Font(name=F, size=10, color=BLUE, underline='single')
    else:
        lc.font = Font(name=F, size=10, color=GREY, italic=True)
        ws.cell(row=row, column=9).fill = PatternFill('solid', fgColor=AMBER)

    pc = ws.cell(row=row, column=10)
    if r['permalink']:
        pc.hyperlink = r['permalink']
        pc.font = Font(name=F, size=10, color=BLUE, underline='single')
    pc.alignment = Alignment(horizontal='center', vertical='top')

    if i % 2 == 1:
        for c in range(1, 11):
            if not (c == 9 and not ok):
                ws.cell(row=row, column=c).fill = PatternFill('solid', fgColor=BAND)

last = HEADER_ROW + len(catalog)
ws.auto_filter.ref = f'A{HEADER_ROW}:J{last}'
ws.freeze_panes = f'A{HEADER_ROW + 1}'

# ─────────────────────────── ชีต 2: สรุปตามหมวดหมู่ ───────────────────────────
s2 = wb.create_sheet('สรุปตามหมวดหมู่')
s2['A1'] = 'จำนวนผลงานแยกตามหมวดหมู่'
s2['A1'].font = Font(name=F, size=14, bold=True, color=NAVY)

cats, seen = [], set()
for r in catalog:
    if r['catName'] not in seen:
        seen.add(r['catName']); cats.append(r['catName'])
cats.sort()

for c, h in enumerate(['หมวดหมู่', 'จำนวนผลงาน', 'สัดส่วน', 'รีแอ็กชันรวม'], 1):
    cell = s2.cell(row=3, column=c, value=h)
    cell.font = Font(name=F, size=10, bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor=NAVY)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = box
for c, w in enumerate([34, 14, 12, 16], 1):
    s2.column_dimensions[get_column_letter(c)].width = w

DATA = f"'แคตตาล็อกผลงาน'!$B${HEADER_ROW+1}:$B${last}"
REACT = f"'แคตตาล็อกผลงาน'!$H${HEADER_ROW+1}:$H${last}"
for i, name in enumerate(cats):
    row = 4 + i
    s2.cell(row=row, column=1, value=name).font = Font(name=F, size=10)
    s2.cell(row=row, column=2, value=f'=COUNTIF({DATA},$A{row})')
    s2.cell(row=row, column=3, value=f'=IFERROR(B{row}/$B${4+len(cats)},0)')
    s2.cell(row=row, column=4, value=f'=SUMIF({DATA},$A{row},{REACT})')
    for c in range(1, 5):
        cell = s2.cell(row=row, column=c)
        cell.border = box
        if c > 1:
            cell.font = Font(name=F, size=10)
            cell.alignment = Alignment(horizontal='center')
    s2.cell(row=row, column=3).number_format = '0.0%'
    s2.cell(row=row, column=2).number_format = '#,##0'
    s2.cell(row=row, column=4).number_format = '#,##0'

tot = 4 + len(cats)
s2.cell(row=tot, column=1, value='รวมทั้งหมด').font = Font(name=F, size=10, bold=True)
s2.cell(row=tot, column=2, value=f'=SUM(B4:B{tot-1})')
s2.cell(row=tot, column=3, value=f'=IFERROR(B{tot}/$B${tot},0)')
s2.cell(row=tot, column=4, value=f'=SUM(D4:D{tot-1})')
for c in range(1, 5):
    cell = s2.cell(row=tot, column=c)
    cell.font = Font(name=F, size=10, bold=True)
    cell.fill = PatternFill('solid', fgColor=BAND)
    cell.border = box
    if c > 1:
        cell.alignment = Alignment(horizontal='center')
s2.cell(row=tot, column=2).number_format = '#,##0'
s2.cell(row=tot, column=3).number_format = '0.0%'
s2.cell(row=tot, column=4).number_format = '#,##0'

s2.cell(row=tot + 2, column=1,
        value='ตัวเลขทั้งหมดคำนวณด้วยสูตรจากชีต "แคตตาล็อกผลงาน" — แก้ข้อมูลในชีตนั้นแล้วตารางนี้อัปเดตเอง'
       ).font = Font(name=F, size=9, italic=True, color=GREY)

# ─────────────────────────── ชีต 3: วิธีทำและข้อจำกัด ───────────────────────────
s3 = wb.create_sheet('วิธีทำและข้อจำกัด')
s3.column_dimensions['A'].width = 30
s3.column_dimensions['B'].width = 96

broken = [r for r in catalog if r['linkStatus'] != 'ใช้ได้']
BLOCKS = [
    ('ที่มาของข้อมูล', ''),
    ('โพสต์ต้นทาง', POST),
    ('วิธีเก็บข้อมูล', 'ดึงคอมเมนต์จาก DOM ของหน้าโพสต์ด้วยสคริปต์ tools/fb-comment-extractor.js '
                      'รันใน DevTools Console ของเบราว์เซอร์ที่ล็อกอิน Facebook '
                      '(Facebook ปิด Groups API สำหรับอ่านคอมเมนต์ตั้งแต่ปี 2020 จึงเรียกผ่าน API ไม่ได้)'),
    ('จำนวนไฟล์ที่นำมารวม', 'ผลการรัน 3 ครั้ง (470 / 859 / 1,339 รายการ) นำมารวมและตัดซ้ำด้วย commentId '
                             'เมื่อคอมเมนต์เดียวกันปรากฏหลายครั้ง เลือกเวอร์ชันที่ข้อความยาวที่สุด'),
    ('คอมเมนต์หลังตัดซ้ำ', '1,329 คอมเมนต์ (คอมเมนต์หลัก 243 + ตอบกลับ 1,086)'),
    ('เกณฑ์คัดเข้าแคตตาล็อก', 'คอมเมนต์ที่เจ้าของนำเสนอผลงานของตัวเองพร้อมลิงก์ '
                              'รวมคอมเมนต์หลัก คอมเมนต์ที่เจ้าของตอบเสริมของตัวเอง '
                              'และผลงาน 3 ชิ้นที่ปรากฏเฉพาะในคอมเมนต์ตอบกลับคนอื่น'),
    ('การรวมและการแยกรายการ', 'คอมเมนต์หลักกับคอมเมนต์เสริมของเจ้าของคนเดียวกันถูกยุบเป็นผลงานเดียว '
                              'ส่วนคอมเมนต์ที่นำเสนอผลงานหลายชิ้นที่ทำคนละเรื่อง ถูกแยกเป็นคนละรายการ'),
    ('จำนวนผลงานสุดท้าย', f'{len(catalog)} รายการ'),
    ('', ''),
    ('ข้อจำกัดที่ต้องรู้', ''),
    ('ลิงก์ที่ใช้ไม่ได้', f'{len(broken)} จาก {len(catalog)} รายการ — Facebook ตัดข้อความลิงก์ยาวด้วย "..." '
                          'และสคริปต์เก็บได้เฉพาะข้อความที่แสดง ไม่ได้เก็บ href จริงของแท็กลิงก์ '
                          'ลิงก์ App Store กู้คืนได้จากเลข app id แต่ Play Store และ Chrome Web Store กู้ไม่ได้ '
                          'รายการเหล่านี้ทำเครื่องหมายไว้ในคอลัมน์ "สถานะลิงก์"'),
    ('ยอดรีแอ็กชัน', 'เป็นค่า ณ เวลาที่ดึงข้อมูล ไม่ใช่ค่าปัจจุบัน และคอมเมนต์ที่ไม่มีรีแอ็กชันบันทึกเป็น 0'),
    ('ความครบถ้วน', 'คอมเมนต์ที่ Facebook ซ่อนหรือกรองเป็นสแปมจะไม่อยู่ใน DOM จึงไม่ถูกเก็บ '
                    'ผลงานบางชิ้นอาจตกหล่นถ้าเจ้าของโพสต์ไว้ในคอมเมนต์ที่ถูกซ่อน'),
    ('หมวดหมู่', 'จัดด้วยการอ่านคำอธิบายของเจ้าของผลงานทีละรายการ ไม่ได้ใช้คีย์เวิร์ดอัตโนมัติ '
                 'ผลงานที่คร่อมหลายหมวดถูกจัดตามการใช้งานหลักที่เจ้าของเน้น'),
    ('รายละเอียดผลงาน', 'สรุปจากข้อความที่เจ้าของเขียนเท่านั้น ไม่ได้เข้าไปตรวจสอบเว็บหรือแอปจริง '
                        'เพราะสภาพแวดล้อมที่ประมวลผลไม่มีสิทธิ์ออกอินเทอร์เน็ต'),
    ('ผลงานที่ไม่ระบุรายละเอียด', 'บางรายการเจ้าของแนบแต่ลิงก์โดยไม่อธิบาย จัดไว้ในหมวด "รวมผลงาน & อื่น ๆ" '
                                   'และระบุไว้ในช่องรายละเอียดตามจริง'),
]

r = 1
for label, text in BLOCKS:
    if label and not text:
        s3.cell(row=r, column=1, value=label).font = Font(name=F, size=12, bold=True, color=NAVY)
        r += 1
        continue
    if not label:
        r += 1
        continue
    a = s3.cell(row=r, column=1, value=label)
    a.font = Font(name=F, size=10, bold=True)
    a.alignment = Alignment(vertical='top')
    b = s3.cell(row=r, column=2, value=text)
    b.font = Font(name=F, size=10)
    b.alignment = Alignment(vertical='top', wrap_text=True)
    if text.startswith('http'):
        b.hyperlink = text
        b.font = Font(name=F, size=10, color=BLUE, underline='single')
    s3.row_dimensions[r].height = max(15, 13 * (len(text) // 90 + 1))
    r += 1

r += 1
s3.cell(row=r, column=1, value='รายการที่ลิงก์ใช้ไม่ได้').font = Font(name=F, size=12, bold=True, color=NAVY)
r += 1
for br in broken:
    s3.cell(row=r, column=1, value=f"ลำดับ {br['no']}").font = Font(name=F, size=10, bold=True)
    c = s3.cell(row=r, column=2, value=f"{br['name']} — {br['url']}")
    c.font = Font(name=F, size=10)
    c.fill = PatternFill('solid', fgColor=AMBER)
    r += 1

# บังคับให้ Excel/Sheets คำนวณสูตรใหม่ทันทีที่เปิดไฟล์
# openpyxl เขียนสูตรโดยไม่มีค่าที่ cache ไว้ ถ้าไม่ตั้งค่านี้บางโปรแกรมจะแสดงเป็นช่องว่าง
wb.calculation.fullCalcOnLoad = True
wb.save('ผลงาน-AI-จากคอมเมนต์-Facebook.xlsx')
print('saved | ผลงาน:', len(catalog), '| ลิงก์เสีย:', len(broken), '| หมวด:', len(cats))
